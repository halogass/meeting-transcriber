"""Transcript exports: TXT, Markdown, SRT, DOCX."""

from __future__ import annotations

import io


def ts(seconds: float, srt: bool = False) -> str:
    ms = int(round(seconds * 1000))
    h, rem = divmod(ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    if srt:
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _name(meta: dict, label: str) -> str:
    return (meta.get("speakers") or {}).get(label) or label


def _header(meta: dict) -> list[str]:
    return [meta["title"], f'{meta["created_at"][:16].replace("T", " ")}  ·  {ts(meta.get("duration") or 0)}']


def to_txt(meta: dict, segments: list[dict]) -> str:
    lines = _header(meta) + [""]
    for s in segments:
        lines.append(f'[{ts(s["start"])}] {_name(meta, s["speaker"])}: {s["text"]}')
    return "\n".join(lines) + "\n"


def to_md(meta: dict, segments: list[dict], summary: str | None = None) -> str:
    title, sub = _header(meta)
    out = [f"# {title}", "", f"_{sub}_", ""]
    if summary:
        out += ["## Summary", "", summary.strip(), "", "## Transcript", ""]
    for s in segments:
        out.append(f'**{_name(meta, s["speaker"])}** `{ts(s["start"])}`  ')
        out.append(s["text"])
        out.append("")
    return "\n".join(out)


def to_srt(meta: dict, segments: list[dict]) -> str:
    out = []
    for i, s in enumerate(segments, 1):
        out += [str(i), f'{ts(s["start"], True)} --> {ts(s["end"], True)}', f'{_name(meta, s["speaker"])}: {s["text"]}', ""]
    return "\n".join(out)


def to_docx(meta: dict, segments: list[dict], summary: str | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    title, sub = _header(meta)
    doc.add_heading(title, level=0)
    p = doc.add_paragraph(sub)
    p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if summary:
        doc.add_heading("Summary", level=1)
        for line in summary.strip().splitlines():
            line = line.rstrip()
            if not line:
                continue
            if line.startswith("#"):
                doc.add_heading(line.lstrip("#").strip(), level=2)
            elif line.lstrip().startswith(("- ", "* ")):
                doc.add_paragraph(_strip_md(line.lstrip()[2:]), style="List Bullet")
            else:
                doc.add_paragraph(_strip_md(line))
        doc.add_heading("Transcript", level=1)

    for s in segments:
        para = doc.add_paragraph()
        who = para.add_run(_name(meta, s["speaker"]))
        who.bold = True
        when = para.add_run(f'  {ts(s["start"])}')
        when.font.size = Pt(8)
        when.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        para.add_run("\n" + s["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _strip_md(text: str) -> str:
    return text.replace("**", "").replace("__", "").replace("`", "")
